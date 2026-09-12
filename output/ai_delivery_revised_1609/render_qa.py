from pathlib import Path
import pypdfium2 as pdfium
from docx import Document
import sys
sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).resolve().parent
for name in ['程序使用说明','培训记录']:
    pdf=pdfium.PdfDocument(str(root/'final'/f'{name}.pdf'))
    out=root/'qa'/name
    out.mkdir(exist_ok=True)
    texts=[]
    for i in range(len(pdf)):
        page=pdf[i]
        texts.append(page.get_textpage().get_text_range())
        page.render(scale=1.5).to_pil().save(out/f'page-{i+1}.png')
    (out/'text.txt').write_text('\n'.join(texts),encoding='utf-8')
    print(name,len(pdf),'pages')
print([p.text for p in Document(root/'final'/'程序使用说明.docx').paragraphs if p.text.startswith('ai_model')])
