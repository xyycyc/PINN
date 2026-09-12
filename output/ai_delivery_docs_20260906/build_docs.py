from pathlib import Path
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'final'
OUT.mkdir(parents=True, exist_ok=True)
SOURCE = Path(r'D:\Desktop\待交付')
NAMES = ['01_AI模块程序使用说明', '02_AI模块维护手册', '03_AI模块程序调试维修方法', '04_AI模块培训记录']
for ext in ['docx','pdf']:
    shutil.copy2(SOURCE / ('技术文档_ai模块.'+ext), OUT / (NAMES[0]+'.'+ext))

def font(run, size=11, bold=False, family='宋体'):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr()
    rf=rpr.rFonts
    if rf is None:
        rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for k in ['ascii','hAnsi','eastAsia']:
        rf.set(qn('w:'+k),family)

def create(title):
    d=Document()
    for node in list(d.styles.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
    for node in list(d.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
    s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2);s.bottom_margin=Cm(1.8)
    s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
    s.footer_distance=Cm(.8)
    for name,size in [('Normal',11),('Title',22),('Heading 1',15),('Heading 2',12)]:
        st=d.styles[name]
        st.font.name='宋体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.18
    d.styles['Normal'].paragraph_format.widow_control=True
    d.core_properties.title=title
    d.core_properties.subject='超声波全波形温度场重构理论及方法研究 模块三'
    d.core_properties.author=''
    p=d.add_paragraph(style='Title');font(p.add_run(title),22,True)
    p=d.add_paragraph('超声波全波形温度场重构理论及方法研究')
    font(p.runs[0],10)
    p=d.add_paragraph('适用范围  基于人工智能的温度场反演方法模块')
    font(p.runs[0],10)
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('第 '),9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    font(p.add_run(' 页'),9)
    return d

def para(d,text):
    p=d.add_paragraph();font(p.add_run(text));return p

def h(d,text,level=1):
    p=d.add_paragraph(text,style='Heading '+str(level))
    for r in p.runs:font(r,15 if level==1 else 12,True)
    return p

def page(d,title):
    d.add_page_break();h(d,title)

def steps(d,items):
    for i,t in enumerate(items,1):para(d,f'{i}. {t}')

def code(d,lines):
    for t in lines:
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.line_spacing=1.0
        font(p.add_run(t),9,False,'Consolas')

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
    t._tbl.tblPr.append(margins)
    for idx,values in enumerate([headers]+rows):
        cells=t.rows[0].cells if idx==0 else t.add_row().cells
        for c,w,v in zip(cells,widths,values):
            c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
            font(p.add_run(v),10,idx==0)
            if idx==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');c._tc.get_or_add_tcPr().append(shade)
        trpr=t.rows[idx]._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def field(d,label,lines=1):
    para(d,label)
    for _ in range(lines):
        p=d.add_paragraph('__________________________________________________________________')
        p.paragraph_format.space_after=Pt(9)
        for r in p.runs:font(r,10)

# 维护手册
d=create('AI模块维护手册')
para(d,'本手册规定 AI 温度场反演模块的运行检查、数据与模型备份、版本更换、恢复验证和维护记录要求，供运行管理人员和软件维护人员使用。执行维护前应停止有关任务并保存日志，维护后以原有验证样例复核运行结果。')
h(d,'1 适用系统与配套资料')
para(d,'程序使用方法见配套《AI模块程序使用说明》（原《技术文档_ai模块》）；故障定位与修复步骤见《AI模块程序调试维修方法》；人员操作培训使用《AI模块培训记录》。本手册仅覆盖 AI 模块及其 Python 与 Fortran 调用环境。')
table(d,['项目','统一运行约定'],[
('系统与环境','Windows；Python 3.13 或更高版本；Fortran 编译器不低于 14.2.0。PyTorch GPU 环境与本机 CUDA、显卡驱动匹配，GUI 使用 tkinter。'),
('模型调用','按材料分别保存完整 checkpoint，通过材料路由选择模型；材料类别与模型、数据清单保持对应。'),
('数据与结果','输入根目录默认为 ai_model/database，输出根目录默认为 ai_model/result；自定义路径以实际运行配置为准。'),
('节点与单位','二维输出对应 10000 个固定物理节点；一维结果提取归一化 x=0.5 中心线；温度单位 K，坐标单位 m。')],[3.1,13.5])
h(d,'2 维护职责与时机')
para(d,'运行人员负责任务前后检查、运行日志保存和异常登记；软件维护人员负责环境、配置、模型与依赖问题的处理；结果复核人员核对维护后的输出与原基线。实际人员由使用单位指定并记入维护记录。')
para(d,'以下维护频次用于安排日常工作：每次运行前检查环境与路径，任务完成后核对产物；每次建库、训练、增量训练或版本更换后执行一次完整备份；定期检查备份可读性和存储空间，周期由使用单位根据运行频率确定。')

page(d,'3 日常检查与资料管理')
h(d,'3.1 任务执行前后检查',2)
steps(d,[
'确认使用交付版本和对应 Python 环境，检查 GPU 可用性、输入目录可读性、输出目录可写性及剩余存储空间。',
'在“路径设置”中核对数据根目录、结果根目录和工作目录；核对实际使用的数据清单、材料规则和模型路由。命令行应在 ai_model 的上一级目录执行。',
'建库后先使用“校验数据”检查清单、划分、材料、温度范围及采样一致性；发现异常时保留报告，处理后重新校验。',
'运行时查看日志与任务状态。任务完成后确认 checkpoint、训练报告或预测结果已经生成，并核对样本数、节点数、材料身份和单位。',
'保存本次配置、日志及结果路径；对异常退出或中断任务标注状态，复核文件完整性后再决定是否使用该次产物。'])
h(d,'3.2 备份对象',2)
table(d,['对象','应成套保存的内容'],[
('程序与环境','交付代码、requirements.txt、实际 Python 与依赖版本、Fortran 源码和已验证的动态库、构建命令及日志。'),
('原始和处理数据','原始波形及温度场；material_collection.json、各材料 manifest、训练验证测试清单、split_config.json、mesh_audit.json、sampling_index.npz 及被引用的数据文件。'),
('模型与配置','对应的所有材料 checkpoint、材料路由文件、database/rule 下的规则登记，以及 window/settings.json（如已生成）和实际运行参数。'),
('报告与预测','训练报告、运行日志、predictions.csv/.npz、metadata.json、metrics.json，以及实际生成的一维结果、对比图和独立测点评价记录。')],[3.1,13.5])
para(d,'备份按任务或版本保存到独立目录，登记时间、原路径、备份路径、文件数量和校验值。复制结束后核对关键文件能够读取；多材料路由和其引用的所有模型必须作为一个整体保存。')

page(d,'4 备份恢复与版本更换')
h(d,'4.1 恢复数据和模型',2)
steps(d,[
'停止相关任务，记录故障时的目录、参数和日志；对当前状态另行留存，避免直接覆盖仍可追溯的资料。',
'选择最近一次经过验证且数据、模型、路由成套的备份，复制到独立恢复目录，并核对文件数量或校验值。',
'检查 manifest、材料集合及模型路由中引用的路径。优先恢复原目录结构；目录必须变更时，通过正常建库、配置或登记流程更新引用并逐项核对，保留原文件。',
'在 GUI 中重新选择恢复目录、清单和对应模型；执行数据校验，再用保留的验证样例进行预测。',
'核对输出文件、样本与节点数量、单位、材料路由和关键结果。将恢复验证记录与备份编号关联，确认可用后恢复日常运行。'])
h(d,'4.2 更新和回退',2)
para(d,'更换程序、Python、依赖、Fortran 动态库或模型前，先保存当前可用版本及环境信息。新版本在独立目录验证，沿用固定的验证输入和评价方法，避免在原数据目录内反复覆盖。')
para(d,'更新后依次检查程序启动、数据校验、模型加载、材料路由、温度场预测和结果导出。涉及训练或增量训练的更新，还应验证新模型保存与再次加载。出现未解决异常时，停止使用新版本，恢复原程序、环境配置和成套模型，再重复验证。')
para(d,'增量训练的基础 checkpoint 与新 checkpoint 分别保留，并保存新增数据清单和参数。选择新模型前先验证对应输出。不得通过手工修改节点数、采样指纹或温度统计量绕过一致性检查。')
h(d,'4.3 项目清理',2)
para(d,'清理前完成备份并核对数据根目录、结果根目录及候选任务。在“项目清理”中按任务规则确认拟处理的模型、报告和预测结果，核对路径后执行界面确认。材料路由及其引用模型按完整组合处理；清理后检查保留任务能正常选择和预测。')

page(d,'5 维护后验证与维护记录')
para(d,'维护验证使用固定的已留存样例，记录输入、模型、运行环境、输出路径和判断依据。一般维护先核对流程与文件完整性；涉及精度或效率的变更，再按交付测试方法复核相应指标。')
para(d,'温度精度按五个代表性测点先分别求参考平均温度和 AI 平均温度，再计算两者的相对误差。中心线误差、全场 MAE/RMSE 和局部区域误差用于分析。无真实温度标签的输入仅验证推理流程，精度评价应结合独立参考测温记录。GPU 效率复核采用与交付测试一致的计时口径。')
table(d,['验证项目','记录要求'],[
('环境与启动','记录 Python、PyTorch、GPU 与 Fortran 环境，检查 GUI 启动及所用命令入口。'),
('数据与模型','记录清单、材料路由及 checkpoint 路径；数据校验无未处理错误，模型和采样信息匹配。'),
('结果与追溯','记录预测文件、样本数和节点数；核对单位、材料和一维或二维输出；保存日志及必要的对比结果。'),
('处置结论','记录已恢复、需继续处理或已回退，并由执行人与复核人填写实际结论。')],[3.1,13.5])
h(d,'5.1 维护记录表',2)
table(d,['记录字段','填写内容'],[
('维护编号和时间','________________________________________________'),
('执行人和复核人','________________________________________________'),
('任务与维护原因','________________________________________________'),
('版本和环境变化','________________________________________________'),
('备份位置及校验信息','________________________________________________'),
('处理步骤及恢复位置','________________________________________________'),
('验证样例及结果位置','________________________________________________'),
('遗留问题及结论','________________________________________________'),
('签字和日期','________________________________________________')],[4.3,12.3])
d.save(OUT/(NAMES[1]+'.docx'))

# 调试维修方法
d=create('AI模块程序调试维修方法')
para(d,'本文件用于定位并处理 AI 模块在环境配置、数据建库、模型训练、预测及增量训练中的软件故障。处理过程遵循“保存现场、复现问题、检查原因、实施修复、验证恢复”的顺序，原始输入、已验证模型及历史结果应保持可追溯。')
h(d,'1 调试范围与统一约定')
para(d,'运行环境采用 Windows、Python 3.13 或更高版本、Fortran 编译器 14.2.0 或更高版本；PyTorch GPU 环境应与 CUDA 和驱动匹配。日常操作沿用《AI模块程序使用说明》，备份和版本回退按《AI模块维护手册》执行。')
para(d,'交付流程采用分材料完整 checkpoint 与材料路由，二维结果为 10000 个固定物理节点，一维结果为二维场归一化 x=0.5 中心线。排查时应保留同一任务的数据清单、采样定义、预处理参数和模型对应关系。')
h(d,'2 先保存故障现场')
steps(d,[
'保存 GUI 运行日志或完整命令行输出，记录首个有效异常及其前后的上下文。',
'登记出错时间、操作页面或子命令、数据与结果根目录、清单路径、材料路由、checkpoint 和主要参数。',
'确认任务已结束或通过“停止”退出，再处理文件。重试时使用独立任务名或输出目录，避免混入上次中断的产物。',
'使用同一输入复现问题；一次只调整一个相关因素，记录修改前后的表现。先备份再变更程序、依赖、配置和模型。'])
h(d,'2.1 环境信息采集',2)
para(d,'在 ai_model 的上一级目录、使用启动 GUI 的同一 Python 环境执行以下命令并保留输出。')
code(d,['python --version','python -m pip show torch numpy pandas scipy Pillow','python -c "import sys; print(sys.executable)"','python -c "import torch; print(torch.__version__); print(torch.cuda.is_available(), torch.version.cuda)"','python -c "import tkinter; print(tkinter.TkVersion)"','gfortran --version'])

page(d,'3 启动与运行环境故障')
h(d,'3.1 找不到程序包或依赖',2)
para(d,'现象：出现 No module named ai_model，或缺少 torch、numpy、PIL 等依赖。首先检查当前目录是否为 ai_model 的上一级，再核对命令行与 GUI 使用的 Python 是否一致。')
para(d,'处理：切换到实际安装目录的上一级，用同一解释器安装交付依赖。requirements.txt 位于 ai_model 目录内，以下命令据此给出相对路径。GPU 版 PyTorch 应使用与交付环境匹配的安装包。')
code(d,[r'python -m pip install -r .\ai_model\requirements.txt','python -m ai_model --help','python -m ai_model.window'])
para(d,'恢复验证：命令帮助可显示，GUI 可启动，运行日志中无缺少依赖的错误。若 tkinter 导入失败，修复该 Python 安装中的 Tcl/Tk 组件后复查。')
h(d,'3.2 GPU 不可用或显存不足',2)
para(d,'现象：CUDA 检查返回 False、设备初始化失败，或训练预测提示显存不足。检查显卡驱动、PyTorch 的 CUDA 支持及当前显存占用，确认实际使用的是目标 Python 环境。')
para(d,'处理：先结束无关 GPU 任务，并避免同时启动多个训练任务；按交付版本修复相匹配的驱动或 PyTorch 环境。需要缩小任务时仅调整界面或命令帮助中支持的参数，并记录变化。可使用 CPU 小样例辅助定位，恢复 GPU 后再验证交付效率。')
para(d,'恢复验证：CUDA 检查通过，原问题样例可以完整执行。CPU 调试耗时不作为 GPU 效率的复核依据。')
h(d,'3.3 中文乱码或目录不可访问',2)
para(d,'中文乱码时先保存原日志，确认 GUI 使用 UTF-8；命令行可执行下列设置后重新运行。目录错误时核对路径、引号、工作目录和读写权限，选择可写的独立结果目录。')
code(d,['$env:PYTHONUTF8 = "1"','$env:PYTHONIOENCODING = "utf-8"'])
para(d,'恢复验证：中文日志可读；运行创建的结果位于预期目录，且路径设置与实际命令一致。')

page(d,'4 数据建库与模型匹配故障')
h(d,'4.1 没有发现样本或波形与温度场配对失败',2)
para(d,'检查所选目录是否为单材料 case 目录的上层或多材料目录的共同上层。标准输入支持 worker 下的 case 和扁平 case 两种布局。核对配置文件、超声波形和热力温度 CSV 是否属于同一 case，避免用不同任务文件补齐缺失项。')
para(d,'处理：恢复缺失文件、纠正目录或列配置后，在新数据集目录重新建库；核对 case 数量、配对记录与报错信息，再执行“校验数据”。对于独立实验波形，按使用说明通过“外部测试导入”挂接到指定材料测试清单。')
h(d,'4.2 数据格式或采样信息异常',2)
para(d,'检查温度场 CSV 的 node_id、x、y、T、thermal_material_id、dup_target_surface 等字段及单位；核对波形时间列、信号列和实际读取范围。检查 sampling_index.npz、网格审计和各份清单是否来自同一次建库。')
para(d,'处理：根据原始数据纠正字段或单位，保持接触界面两侧节点的独立身份；从完整输入重新建库并校验。不得仅更改数组长度、删除重复坐标节点或改写指纹来消除报错。')
h(d,'4.3 模型文件不存在或一致性检查失败',2)
para(d,'检查 checkpoint 路径及材料路由中每个模型引用，核对模型类型、节点数、采样指纹、波形长度、预处理参数和温度统计量。找不到模型时从同一任务备份恢复完整文件组；模型与数据不匹配时选择同一建库任务对应的模型，或用正确数据重新训练。')
para(d,'不得以其他材料 checkpoint 替代缺失模型。材料文件夹代表样本材料类别，温度场中的内部材料组分用于节点身份，两者用途不同。修复后逐材料检查路由能够选择对应模型。')
h(d,'4.4 数据集划分或校验不通过',2)
para(d,'检查各材料训练、验证和测试比例及清单引用，依据界面提示修正无效配置。固定训练集及其统计量，保留划分参数和种子。独立外部实验测试数据按测试用途管理，避免误并入训练或验证集。重新建库后记录变化，并再次校验。')

page(d,'5 运行结果与增量训练故障')
h(d,'5.1 训练中断或没有预期模型',2)
para(d,'从日志区分正常结束、手动停止、早停与异常退出。检查数据加载、输出目录、显存和磁盘空间，核对训练报告及 checkpoint 的实际保存位置。先消除具体异常，再以独立任务名重试；不得把中断文件直接登记为已验证模型。')
para(d,'恢复验证：日志正常结束，训练报告可读取，模型能重新加载并对保留样例生成预测。多材料任务需检查全部材料模型及路由文件。')
h(d,'5.2 预测图或结果文件缺失',2)
para(d,'先核对任务是否成功、实际输出根目录以及预测名称。检查是否启用绘图，是否选择一维或二维输出。二维绘图对应 point_field_compare.png；一维输出对应 axis_predictions.csv/.npz，启用绘图时生成 axis_compare.png。')
para(d,'恢复验证：按本次选项检查实际应生成的文件，使用独立输出目录重试。只生成原始数值结果且未启用绘图时，无对比图不应误判为预测失败。')
h(d,'5.3 温度结果异常或指标发生变化',2)
para(d,'依次核对样本材料与路由、波形列和裁剪长度、预处理、模型和采样定义、温度单位及参考数据。用已留存的固定样例对比维护前后结果，定位输入或配置变化，不能只根据图像外观修改预测值。')
para(d,'五测点验收精度按五点参考温度均值与五点 AI 温度均值计算相对误差；MAE、RMSE、中心线和区域误差用于补充分析。独立实验波形的精度复核需有参考测温记录，不能将缺乏真实空间标签的场指标作为真实全场精度。')
h(d,'5.4 增量训练失败或新模型表现下降',2)
para(d,'检查基础 checkpoint、新增清单、材料身份、采样定义和预处理的一致性，先确认基础模型可正常预测。修正输入配置后再执行增量训练，保留基础模型和独立的新模型。新模型未通过验证时，重新选择已验证基础模型，并记录本次失败及差异。')
h(d,'5.5 界面无响应或停止后的文件不完整',2)
para(d,'先查看运行日志和任务状态，区分正在计算与已退出。需停止时使用界面“停止”，确认子进程退出后再更改输入或清理输出。重新打开 GUI 后核对路径与任务选择，并在新输出目录验证；保存异常现场供进一步排查。')

page(d,'6 Fortran 调用及修复后的确认')
h(d,'6.1 动态库缺失或加载失败',2)
para(d,'Fortran library not found 表示调用所需动态库未在预期位置找到。交付 Windows 程序使用 ai_model/fortran/ai_numeric.dll。核对交付文件、编译器、运行库与进程架构，优先从同一已验证交付版本恢复。')
para(d,'需要重新构建时，先备份原动态库，在 ai_model 目录下执行交付脚本；不要从不相关版本复制动态库。以下命令中的目录必须切换为实际安装位置。')
code(d,['gfortran --version',r'powershell -ExecutionPolicy Bypass -File .\fortran\build_fortran.ps1'])
para(d,'在 ai_model 上一级目录用同一 Python 环境检查动态库是否可加载：')
code(d,['python -c "import ctypes; ctypes.CDLL(\'ai_model/fortran/ai_numeric.dll\'); print(\'DLL load OK\')"'])
para(d,'动态库加载成功后，还须运行原先发生错误的实际业务样例，复核数值输出和日志。加载成功本身不代表全部功能与数值结果已通过验证。')
h(d,'6.2 修复完成条件',2)
steps(d,[
'原故障样例能在目标环境正常执行，完整日志没有未处理异常。',
'数据、材料路由和模型匹配，输出样本数、节点数、单位及文件完整性符合本次任务配置。',
'必要的精度或效率复核沿用交付测试方法；不能复核的项目明确记录为待验证。',
'保存修改项、前后版本、验证结果、执行人与复核人；有未解决问题时记录临时措施或按维护手册回退。'])
h(d,'6.3 故障处理记录',2)
table(d,['字段','填写内容'],[
('编号与发现时间','________________________________________________'),
('现象及完整日志位置','________________________________________________'),
('环境和数据模型路径','________________________________________________'),
('原因及处理步骤','________________________________________________'),
('备份与修改前后版本','________________________________________________'),
('验证样例及结果位置','________________________________________________'),
('遗留问题与处理结论','________________________________________________'),
('执行人和复核人签字','________________________________________________')],[4.3,12.3])
d.save(OUT/(NAMES[2]+'.docx'))

# 培训记录
d=create('AI模块培训记录')
para(d,'本记录用于登记 AI 温度场反演模块的人员培训、操作练习、结果确认和问题处理情况。培训内容按配套使用说明、维护手册及调试维修方法组织；日期、人员、实际完成项目和核查结论由现场据实填写并签认。')
h(d,'1 培训基本信息')
para(d,'项目名称：超声波全波形温度场重构理论及方法研究')
para(d,'培训模块：基于人工智能的温度场反演方法模块')
field(d,'记录编号及培训日期')
field(d,'培训地点及起止时间')
field(d,'培训组织单位及讲师姓名')
field(d,'参训单位及人数')
field(d,'实际培训方式及所用程序版本')
h(d,'1.1 培训环境与教材',2)
para(d,'按使用说明核对 Windows、Python 3.13 或更高版本、Fortran 编译器 14.2.0 或更高版本、PyTorch GPU 环境及 tkinter；实际环境信息填写如下。')
field(d,'实际计算机配置及软件环境信息或清单编号')
para(d,'配套教材：01_AI模块程序使用说明；02_AI模块维护手册；03_AI模块程序调试维修方法。实际使用的文件版本或归档位置：')
field(d,'教材版本或归档位置')

page(d,'2 培训内容与实际完成记录')
para(d,'下列条目作为培训记录项目。完成情况、课时和相关证据由讲师根据实际实施情况填写；未实施项目注明原因，不预填“已完成”或“合格”。')
table(d,['培训项目','内容及操作要点','实际情况与证据'],[
('环境与启动','核对环境；启动 GUI；设置输入输出根目录和工作目录。','________________\n________________'),
('数据库构建','单材料或多材料导入；划分清单；独立实验波形挂接；执行数据校验。','________________\n________________'),
('模型训练','选择正确清单与材料；设置设备和训练参数；检查各材料 checkpoint 和路由。','________________\n________________'),
('预测与结果','选择匹配模型；二维节点场和一维中心线；查看数值文件、对比图及日志。','________________\n________________'),
('评价口径','五测点先求平均再计算相对误差；区分补充分析指标；记录效率测试条件。','________________\n________________'),
('增量训练','选择基础模型及新增数据；保存新模型；验证后使用并保留基础版本。','________________\n________________'),
('维护与恢复','备份程序、清单、采样文件、模型和路由；恢复验证；项目清理前核对。','________________\n________________'),
('故障排查','保存日志；检查入口、GPU、目录、数据和模型匹配；复现与修复验证。','________________\n________________')],[3,8.2,5.4])
h(d,'2.1 操作核查与问题记录',2)
para(d,'以实际样例核查参训人员能否完成启动和路径设置、数据与模型选择、预测及结果查看、日志保存和常见故障定位。结果填写为已掌握、需补训或未核查，并注明依据。')
field(d,'操作样例和结果归档位置')
field(d,'实际核查结论及需补训项目')
field(d,'提出的问题及处理意见')

page(d,'3 参训签到与培训确认')
para(d,'由实际参训人员本人填写或核对，并在签名栏签字；人数超过本页容量时附签到续页。')
table(d,['序号','姓名','单位或部门','岗位','本人签名'],[
(str(i),'','', '', '') for i in range(1,9)
],[1.2,2.8,5.8,3,3.8])
h(d,'3.1 培训结果确认',2)
field(d,'讲师对实际培训内容及操作核查结果的确认',2)
field(d,'参训单位代表意见及遗留事项',2)
para(d,'讲师签字：__________________    日期：__________________')
para(d,'参训单位代表签字：____________    日期：__________________')
para(d,'记录人签字：________________    日期：__________________')
h(d,'3.2 附件登记',2)
para(d,'根据实际情况登记签到续页、课件、操作样例、运行日志、考核或补训记录等附件；涉及照片或影像时填写文件名和归档位置。')
field(d,'附件名称及归档位置',2)
d.save(OUT/(NAMES[3]+'.docx'))
print('\n'.join(str(p) for p in OUT.glob('*')))
