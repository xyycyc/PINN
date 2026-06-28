from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "ai反演.docx"
OUTPUT = ROOT / "修改.docx"


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_font(run, name: str = "宋体", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        set_font(run, "黑体", bold=True)


def add_text(document: Document, text: str, style: str = "正文缩进2") -> None:
    paragraph = document.add_paragraph(style=style)
    set_font(paragraph.add_run(text))


def add_numbered_items(document: Document, items: list[str]) -> None:
    # The source template uses abstract numbering 0: decimal items formatted as (1), (2), ...
    num = document.part.numbering_part.element.add_num(0)
    for item in items:
        paragraph = document.add_paragraph(style="List Paragraph")
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(num.numId))
        num_pr.append(ilvl)
        num_pr.append(num_id)
        p_pr.append(num_pr)
        set_font(paragraph.add_run(item))


def add_test_section(
    document: Document,
    title: str,
    scene: list[str],
    conditions: list[str],
    process: list[str],
) -> None:
    add_heading(document, title, 2)
    add_heading(document, "测试场景", 3)
    for text in scene:
        add_text(document, text)
    add_heading(document, "测试条件", 3)
    add_numbered_items(document, conditions)
    add_heading(document, "测试流程", 3)
    add_numbered_items(document, process)


def build() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    clear_body(document)

    add_heading(document, "AI 温度场反演可视化功能测试大纲（修订版）", 1)
    add_text(
        document,
        "本大纲仅覆盖 ai_model.window 当前图形化操作面板中实际展示的 AI 功能，包括路径配置、数据处理、模型训练、数据校验、预测对比、增量训练、一键演示和项目清理。未在当前主界面展示的功能不纳入本大纲。",
        "Normal",
    )
    add_text(
        document,
        "各功能测试统一采用“测试场景、测试条件、测试流程”结构。测试操作原则上通过中文 GUI 完成，并结合运行日志和生成文件核查结果。",
        "Normal",
    )

    add_test_section(
        document,
        "一、可视化面板启动与路径设置",
        [
            "验证用户能够启动 AI 温度场反演图形化操作面板，并通过“路径设置”页统一配置 Python 解释器、子进程工作目录、输入数据根目录和结果输出根目录。",
        ],
        [
            "运行环境为 Windows，Python 版本为 3.10 或以上，并安装 requirements.txt 中列出的依赖。",
            "通过 python -m ai_model.window 启动界面；GUI 依赖 tkinter，图片预览相关能力依赖 Pillow。",
            "默认输入根目录为 database，默认输出根目录为 result；相对路径以 ai_model 包目录为基准解析。",
        ],
        [
            "启动 GUI，确认主窗口、主流程页签、运行日志面板、状态栏和停止任务功能显示正常。",
            "在“路径设置”页分别设置 Python 解释器、子进程工作目录、输入根目录和输出根目录并保存。",
            "重新加载配置或重启 GUI，确认设置被写入 settings.json 并能正确恢复。",
            "输入不存在或无效的解释器与目录，确认界面能够提示参数错误，且不会启动错误任务。",
            "启动任一耗时任务后执行停止操作，确认程序先终止子进程，界面状态和日志能够恢复。",
        ],
    )

    add_test_section(
        document,
        "二、构建数据库与数据集划分",
        [
            "验证“构建数据库”页能够将仿真数据、实验 CSV 数据和外部仿真 CSV 数据转换为 AI 训练与预测使用的结构化数据，并按设置划分训练集和测试集。",
        ],
        [
            "准备可读取的实验数据目录或外部仿真数据目录；需要生成仿真数据时设置每种材料的仿真样本数。",
            "材料名称仅使用项目允许的格式。实验材料标签会登记到 database/rule/material.csv，供后续训练和模型选择复用。",
            "启用数据集划分时，测试集比例必须位于 0 和 1 之间，并设置随机种子及实验样本划分策略。",
        ],
        [
            "填写每种材料仿真样本数，按需选择跳过仿真数据。",
            "选择实验数据目录、实验材料、导入数量，以及可选的外部仿真目录、材料和导入数量。",
            "启用数据集划分，设置测试集比例、随机种子、实验样本策略和训练/测试清单文件名。",
            "点击执行并观察日志，确认任务完成且没有路径、格式或材料规则错误。",
            "检查 database 下生成的波形、温度场和 manifest 文件；启用划分时确认 train_manifest.json 与 test_manifest.json 生成。",
            "使用相同数据、策略和随机种子重复执行，核对划分结果可复现；使用非法比例或错误数据目录时确认界面阻止执行或日志明确报错。",
        ],
    )

    add_test_section(
        document,
        "三、数据校验",
        [
            "验证“校验数据”页能够读取指定 manifest 清单，统计仿真样本和实验样本数量，并检查数据是否满足 AI 模型使用要求。",
        ],
        [
            "使用“构建数据库”页生成的 combined_manifest.json，或选择其他有效 manifest 文件。",
            "项目配置的最低样本要求为仿真样本不少于 1000 组、实验样本不少于 20 组。",
        ],
        [
            "选择待校验的 manifest 文件并执行校验。",
            "在日志中核对总样本数、仿真样本数、实验样本数及最低数量判定。",
            "抽查清单中引用的波形文件、温度场文件、材料、维度、稳态/瞬态和温度标签是否存在且可解析。",
            "分别测试缺失文件、非法 JSON、空清单以及样本数量不足的清单，确认校验失败信息明确。",
        ],
    )

    add_test_section(
        document,
        "四、AI 模型训练",
        [
            "验证“训练模型”页能够根据训练清单和规则三元组启动 AI 温度场重构模型训练，并正确保存检查点、训练报告及规则登记信息。",
        ],
        [
            "准备有效的 train_manifest.json，并在路径设置中配置可用的 Python 环境和输出根目录。",
            "选择反演维度 one/two、稳态或瞬态模式，并填写合法材料名称；训练名称用于 checkpoint/report 子目录命名，检查点名称用于权重文件命名。",
            "根据测试需要设置 cuda、cpu 或 auto 设备，normal 或 residual_pinn 训练模式、训练轮数、分支权重及波形预处理参数。",
        ],
        [
            "选择训练清单，填写材料、维度、模式、训练名称和检查点名称。",
            "设置训练设备、训练模式、轮数及相关超参数后点击执行。",
            "观察日志中的数据加载、训练轮次、损失和保存信息，确认无 NaN、Inf、显存不足或文件错误。",
            "检查 result/train/checkpoint/<训练名称>/ 下的检查点和 result/train/report/<训练名称>/ 下的训练报告。",
            "刷新模型规则后，确认新检查点已按材料、维度和模式登记，并能在“预测对比”和“增量训练”页中选择。",
            "使用缺失清单、非法材料名或不可用 CUDA 设备执行负向测试，确认错误信息明确且不会登记无效检查点。",
        ],
    )

    add_test_section(
        document,
        "五、AI 预测对比与性能测试",
        [
            "验证“预测对比”页能够按材料、维度和稳态/瞬态筛选已训练模型，对测试清单执行 AI 推理，并输出温度、温度场、声学参数及性能统计结果。",
        ],
        [
            "准备与所选检查点兼容的 test_manifest.json；检查点应已登记到模型规则表。",
            "项目直接输出 temperature_mae 和 temperature_rmse。若合同要求相对误差不大于 10%，应使用 predictions.csv 中的真值和预测值另行计算，并预先规定零值处理方法。",
            "性能测试可启用基准测速，设置预热样本数和前向运行次数；纯前向延迟与包含数据读取、汇总的端到端延迟应分别判定。",
        ],
        [
            "依次选择材料、维度和模式，确认检查点下拉框只显示匹配的已登记模型。",
            "选择测试清单和检查点，设置预测名称或输出目录；按需启用对比图、抽样温度场图和性能基准。",
            "执行预测并观察日志，确认模型配置从检查点正确恢复，推理过程正常完成。",
            "检查 predictions.npz、predictions.csv 和 metrics.json；启用绘图时检查散点图及 field_compare 目录中的温度场对比图。",
            "核对样本数、样本编号、材料、真值与预测值的一致性，并按材料、维度、模式分别统计误差。",
            "启用 benchmark 时核对 benchmark_latency_ms_per_sample 和吞吐率；如验收要求小于 1 秒/点，应以统一设备、批大小和预热条件重复测试。",
            "测试不匹配检查点、空清单、缺失数据和无对应规则模型的情况，确认界面或日志能够阻止错误推理。",
        ],
    )

    add_test_section(
        document,
        "六、AI 在线增量训练",
        [
            "验证“增量训练”页能够使用新增数据在已有 AI 检查点基础上继续训练，并保存新的权重、增量报告和规则登记记录。",
        ],
        [
            "准备新增实验数据清单，并确保基础检查点已按材料、维度和模式登记。",
            "界面会根据所选检查点恢复训练模式、分支权重和默认在线训练轮数；用户修改后将覆盖检查点中的默认值。",
            "增量权重默认保存到基础检查点同级目录，文件名为训练时间；报告保存到 report/<基础任务名>/Incremental/<训练时间>/。",
        ],
        [
            "选择材料、维度和模式，再从匹配列表中选择基础检查点。",
            "选择新增数据 manifest，按需填写输出文件名和调整训练运行参数。",
            "记录基础检查点哈希和增量训练前的固定测试集指标，然后执行增量训练。",
            "观察日志，确认模型加载、训练循环、权重保存和规则登记正常完成。",
            "检查新检查点、增量训练报告以及 trained_rules.csv 登记信息，并在预测页加载新检查点复测。",
            "对旧测试集和新增测试集分别比较 MAE/RMSE。灾难性遗忘不是 GUI 自动判定项，应按项目验收标准设置允许的旧集性能退化范围。",
        ],
    )

    add_test_section(
        document,
        "七、AI 一键演示",
        [
            "验证“一键演示”页能够以最少参数串联数据构建、训练及相关 AI 主流程，用于环境联通性和基础功能快速检查。",
        ],
        [
            "路径设置正确，Python 依赖完整，并具备可用的输入数据与输出目录。",
            "设置每种材料仿真样本数、实验数据导入数量和训练名称；冒烟测试应使用较小但有效的数据量，正式验收仍应使用规定规模的数据。",
        ],
        [
            "填写演示参数并点击执行。",
            "观察日志，确认数据准备、训练和产物保存按顺序执行，任务状态能够正确更新。",
            "检查生成的 manifest、检查点和训练报告，确认目录与训练名称一致。",
            "中途停止任务，确认 GUI 能终止当前子进程；再次执行时确认已有产物不会导致错误关联。",
        ],
    )

    add_test_section(
        document,
        "八、AI 项目产物清理",
        [
            "验证“项目清理”页能够按材料、维度、模式和指定检查点扫描关联的 AI 权重、训练报告、预测结果及规则登记，并在用户确认后执行清理。",
        ],
        [
            "选择正确的数据根目录和结果根目录，并确保 trained_rules.csv 中存在可选检查点。",
            "清理操作具有破坏性，必须先扫描关联产物并核对删除预览；测试时应先使用可恢复的测试模型和备份数据。",
        ],
        [
            "依次选择材料、维度和模式，再选择目标检查点。",
            "点击“扫描关联产物”，核对预览中的目标检查点、训练报告、预测目录和 CSV 登记项。",
            "取消确认，验证不会删除任何文件；再次执行并确认删除，检查界面提示和日志。",
            "核对目标权重及关联产物已按选择删除，其他材料、规则组合和检查点未受影响。",
            "刷新模型列表，确认被删除检查点不再显示；测试失效路径或缺失文件时，确认清理流程能够给出明确结果而不误删其他产物。",
        ],
    )

    document.core_properties.title = "AI 温度场反演可视化功能测试大纲（修订版）"
    document.core_properties.subject = "ai_model.window 可视化功能测试"
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
