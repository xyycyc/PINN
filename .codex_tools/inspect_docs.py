from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document


def dump_docx(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append({
                "index": index,
                "style": paragraph.style.name if paragraph.style else "",
                "text": text,
            })
    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})
    return {"paragraphs": paragraphs, "tables": tables}


if __name__ == "__main__":
    source = Path(sys.argv[1])
    print(json.dumps(dump_docx(source), ensure_ascii=False, indent=2))
